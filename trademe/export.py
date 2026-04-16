"""Export TradeMe data to Excel spreadsheets."""
import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill(start_color="1A56DB", end_color="1A56DB", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
ALT_FILL = PatternFill(start_color="EBF3FE", end_color="EBF3FE", fill_type="solid")
BORDER = Border(
    bottom=Side(style="thin", color="D1D5DB"),
)


def _style_header_row(ws, row: int, num_cols: int):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_data_row(ws, row: int, num_cols: int):
    fill = ALT_FILL if row % 2 == 0 else None
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        if fill:
            cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER


def _auto_width(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                val = str(cell.value or "")
                max_len = max(max_len, len(val))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


def export_orders_to_excel(orders: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.row_dimensions[1].height = 30

    headers = [
        "Order ID", "Listing ID", "Title", "Buyer Username", "Buyer Email",
        "Buyer Phone", "Quantity", "Price (NZD)", "Shipping Cost", "Total (NZD)",
        "Payment Method", "Payment Status", "Shipping Method", "Status",
        "Delivery Name", "Delivery Address", "Delivery City",
        "Delivery Postcode", "Delivery Country", "Note", "Sale Date",
    ]

    for col_idx, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_idx, value=header)
    _style_header_row(ws, 1, len(headers))

    for row_idx, order in enumerate(orders, 2):
        buyer = order.get("Buyer") or {}
        shipping = order.get("ShippingAddress") or {}
        price = order.get("BuyNowPrice") or order.get("Price") or 0
        ship_cost = order.get("ShippingCost") or 0

        row_data = [
            order.get("OrderId") or order.get("PurchaseId"),
            order.get("ListingId"),
            order.get("Title", ""),
            buyer.get("Nickname") or order.get("BidderNickname", ""),
            buyer.get("Email", ""),
            buyer.get("PhoneNumber", ""),
            order.get("QuantitySold") or order.get("Quantity") or 1,
            price,
            ship_cost,
            round(float(price) + float(ship_cost), 2),
            order.get("PaymentMethod", ""),
            order.get("IsPaid") and "Paid" or "Unpaid",
            order.get("ShippingType") or order.get("ShippingOption", ""),
            order.get("Status", ""),
            shipping.get("Name", ""),
            shipping.get("Address", ""),
            shipping.get("City", ""),
            shipping.get("Postcode", ""),
            shipping.get("Country", ""),
            order.get("NoteToSeller", ""),
            order.get("EndDate") or order.get("SaleDate", ""),
        ]
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)
        _style_data_row(ws, row_idx, len(headers))

    # Freeze header row
    ws.freeze_panes = "A2"
    _auto_width(ws)

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws2.append(["Total Orders", len(orders)])
    total_revenue = sum(
        float(o.get("BuyNowPrice") or o.get("Price") or 0) +
        float(o.get("ShippingCost") or 0)
        for o in orders
    )
    ws2.append(["Total Revenue (NZD)", round(total_revenue, 2)])
    paid = sum(1 for o in orders if o.get("IsPaid"))
    ws2.append(["Paid Orders", paid])
    ws2.append(["Unpaid Orders", len(orders) - paid])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def export_listings_to_excel(listings: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Listings"
    ws.row_dimensions[1].height = 30

    headers = [
        "Listing ID", "Title", "Category", "Buy Now Price", "Reserve",
        "Start Price", "Quantity", "Quantity Sold", "Views", "Watchlist Count",
        "Start Date", "End Date", "Status", "Has Buy Now",
        "Shipping Options", "Promoted",
    ]

    for col_idx, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_idx, value=header)
    _style_header_row(ws, 1, len(headers))

    for row_idx, listing in enumerate(listings, 2):
        row_data = [
            listing.get("ListingId"),
            listing.get("Title", ""),
            listing.get("Category", ""),
            listing.get("BuyNowPrice"),
            listing.get("ReservePrice"),
            listing.get("StartPrice"),
            listing.get("Quantity"),
            listing.get("QuantitySold"),
            listing.get("ViewCount"),
            listing.get("WatcherCount"),
            listing.get("StartDate", ""),
            listing.get("EndDate", ""),
            listing.get("ListingStatus", "Active"),
            listing.get("HasBuyNow") and "Yes" or "No",
            ", ".join(
                s.get("Description", "") for s in (listing.get("ShippingOptions") or [])
            ),
            listing.get("IsPromoted") and "Yes" or "No",
        ]
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)
        _style_data_row(ws, row_idx, len(headers))

    ws.freeze_panes = "A2"
    _auto_width(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
